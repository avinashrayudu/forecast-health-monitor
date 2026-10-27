"""
Compares the recalculated workbook with the Python engine, line by line and week by week.

  python -m sheets.check_xlsx        exits non-zero if any closing-cash week differs by more than 1%

Built by Avinash Rayudu.
"""
from __future__ import annotations

import sys

import numpy as np
from openpyxl import load_workbook

from sheets.build_xlsx import OUT, col
from world.engine import N_WEEKS, forecast
from world.inject_events import load_configs

ROWS = {14: "dtc_payouts", 17: "marketplace_payouts", 21: "wholesale_receipts", 24: "payroll", 25: "inventory",
        26: "freight_duty", 27: "fulfillment_3pl", 28: "other_ap", 29: "card_autopay", 30: "fixed_costs",
        31: "sales_tax", 32: "debt_service", 33: "other_opex"}


def compare():
    fc = forecast(load_configs()["fernwood-goods"])
    ws = load_workbook(OUT, data_only=True)["Forecast"]
    if ws[f"{col(0)}37"].value is None:
        raise SystemExit("workbook has no cached values; open and save it in Excel or LibreOffice first")
    out = {"lines": {}}
    for r, line in ROWS.items():
        sheet = np.array([float(ws[f"{col(w)}{r}"].value or 0) for w in range(N_WEEKS)])
        out["lines"][line] = float(np.max(np.abs(sheet - fc.lines[line])))
    cash = np.array([float(ws[f"{col(w)}37"].value) for w in range(N_WEEKS)])
    out["max_cash_diff_pct"] = float(np.max(np.abs(cash - fc.cash) / np.abs(fc.cash)))
    out["low_point_sheet"] = float(ws["B44"].value)
    out["low_point_code"] = fc.low_point()[1]
    out["draw_sheet"] = float(ws["B47"].value)
    return out


def main():
    r = compare()
    for line, d in r["lines"].items():
        print(f"  {line:<22} max weekly difference ${d:,.2f}")
    print(f"  closing cash: max difference {r['max_cash_diff_pct']:.4%}")
    print(f"  low point: workbook ${r['low_point_sheet']:,.0f}, code ${r['low_point_code']:,.0f}; "
          f"workbook draw ${r['draw_sheet']:,.0f}")
    if r["max_cash_diff_pct"] > 0.01:
        sys.exit(1)


if __name__ == "__main__":
    main()
