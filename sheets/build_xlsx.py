"""
Builds sheets/fernwood_13wk.xlsx: Fernwood's 13-week cash forecast as a live workbook.

Every forecast number is a formula on the Inputs, Calendar and Bills tabs, so a
finance lead can change an assumption (a refund rate, a pay week, the opening
balance) and watch the low point move. The Check tab holds the Python engine's
closing cash for the same inputs and shows the difference week by week.

  python -m sheets.build_xlsx            write the workbook, then recalculate it with LibreOffice if installed
  python -m sheets.check_xlsx            compare a recalculated copy with the code

The committed copy was recalculated with LibreOffice so it opens with values.

Built by Avinash Rayudu. All numbers are synthetic.
"""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from world.engine import BFCM, BFCM_SHARE, HIST_MULT, MONTH_MULT, N_WEEKS, WEEK0, forecast
from world.inject_events import load_configs

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "sheets" / "fernwood_13wk.xlsx"

F = "Arial"
BLUE = Font(name=F, color="0000FF", size=10)
BLACK = Font(name=F, color="000000", size=10)
GREEN = Font(name=F, color="008000", size=10)
BOLD = Font(name=F, bold=True, size=10)
TITLE = Font(name=F, bold=True, size=13)
NOTE = Font(name=F, italic=True, color="666666", size=9)
YELLOW = PatternFill("solid", fgColor="FFFF00")
GREY = PatternFill("solid", fgColor="F2F2F2")
HEAD = PatternFill("solid", fgColor="DDE4EE")
USD = '$#,##0;($#,##0);"-"'
PCT = '0.0%'
THIN = Border(top=Side(style="thin", color="999999"))

FIRST_W, LAST_W = -12, 12              # forecast tab columns: 12 history weeks + 13 forecast weeks
CAL_START = WEEK0 + timedelta(days=7 * FIRST_W)
CAL_DAYS = 7 * (LAST_W - FIRST_W + 1)
CAL_END_ROW = 1 + CAL_DAYS


def col(w):
    """Forecast-tab column letter for week index w."""
    return get_column_letter(2 + (w - FIRST_W))


def pos(expr):
    """INDEX position in the B:Z week range for a week-index expression."""
    return f"({expr})+{1 - FIRST_W}"


def build():
    cfgs = load_configs()
    cfg = cfgs["fernwood-goods"]
    fc = forecast(cfg)
    wb = Workbook()

    # ------------------------------------------------------------------ Inputs
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = "Fernwood Goods: 13-week cash forecast inputs"
    ws["A1"].font = TITLE
    ws["A2"] = ("Synthetic brand. Blue cells are inputs; black cells are formulas. Values come from the corrected "
                "plan P&L (docs/fernwood_pnl.md) and the onboarded ledger's calendars.")
    ws["A2"].font = NOTE
    p = cfg["plan"]
    rows = [
        ("week0", "Forecast week 0 starts (Monday)", WEEK0, "date", "Actuals through Oct 25, 2026"),
        ("base", "Shopify gross sales, average history week", cfg["dtc_weekly_gross"], USD,
         "Plan revenue x DTC mix, scaled to an average Jul-Oct week"),
        ("amz", "Amazon sales, average history week", cfg["amz_weekly_sales"], USD, "Plan revenue x Amazon mix"),
        ("wh", "Wholesale invoices, average history week", cfg["wh_weekly"], USD, "Plan revenue x wholesale mix"),
        ("hist", "History seasonality (avg of Jul-Oct multipliers)", None, "0.0000", "Formula on the month table"),
        ("refund", "Refund rate", cfg["refund_rate"], PCT, "Plan: 6% of revenue"),
        ("refund_jan", "Refund rate, January", cfg["refund_rate_jan"], PCT, "Plan: post-holiday returns"),
        ("fee", "Payment processing fees", cfg["fee_rate"], PCT, "Plan: 3% of revenue"),
        ("cb", "Chargebacks", cfg["chargeback_rate"], PCT, "Default for the portfolio"),
        ("lag", "Share of a week's net paid out the following week", cfg["payout_lag_share"], PCT,
         "About 2 of 7 days"),
        ("amz_fee", "Amazon fees", cfg["amz_fee_rate"], PCT, "Referral and FBA, portfolio default"),
        ("faire_share", "Faire share of wholesale", cfg["retailers"][0]["share"], PCT, "Kickoff call"),
        ("faire_terms", "Faire pays after (weeks)", cfg["retailers"][0]["terms_weeks"], "0", "Kickoff call"),
        ("hud_share", "Hudson Mercantile share of wholesale", cfg["retailers"][1]["share"], PCT, "Kickoff call"),
        ("hud_terms", "Hudson Mercantile pays after (weeks)", cfg["retailers"][1]["terms_weeks"], "0",
         "Net 60 plus processing"),
        ("ad_rate", "Ad spend as share of Shopify sales", cfg["ad_rate"], PCT, "Plan: 22% of DTC; Meta 70%, Google 30%"),
        ("oct_ads", "October ad spend per week (history)", None, USD, "Formula: base x ad rate x Oct multiplier"),
        ("card_other", "Software and other card spend per week", cfg["card_other_weekly"], USD,
         "Plan: $10K a month, part of the $34K fixed"),
        ("payroll", "Payroll per run (biweekly)", cfg["payroll_run"], USD, "Plan: $50K biweekly"),
        ("fixed", "Fixed costs paid from the bank on the 1st", cfg["fixed_monthly"], USD,
         "Plan: $34K a month, less $10K paid on the card"),
        ("tax", "Sales tax remitted on the 20th, share of last month's Shopify sales", cfg["sales_tax_rate"], PCT,
         "Plan"),
        ("other_opex", "Bank fees and other small payments per week", cfg["other_opex_weekly"], USD, "Ledger"),
        ("opening", "Opening cash, Oct 26", cfg["opening_cash"], USD, "Lowered from the first draft's $1.0M"),
        ("buffer", "Cash buffer policy", cfg["cash_buffer"], USD, "Founder's minimum"),
        ("loc_limit", "Credit line limit", cfg["loc_limit"], USD, "Illustrative"),
        ("loc_drawn", "Credit line drawn today", cfg["loc_drawn"], USD, "Illustrative"),
        ("apr", "Credit line APR", cfg["loc_apr"], PCT, "Illustrative"),
        ("draw_inc", "Draws are made in multiples of", 25000, USD, "actions/vendor_rules.yaml"),
        ("bfcm_start", "BFCM week starts", BFCM[0], "date", ""),
        ("bfcm_end", "BFCM week ends", BFCM[1], "date", ""),
        ("bfcm_share", "Share of November sales in BFCM week", BFCM_SHARE, PCT, "Seasonality assumption"),
    ]
    N = {}
    ws["A4"], ws["B4"], ws["C4"] = "Input", "Value", "Source / note"
    for c in "ABC":
        ws[f"{c}4"].font, ws[f"{c}4"].fill = BOLD, HEAD
    r = 5
    for key, label, val, fmt, note in rows:
        ws.cell(r, 1, label).font = BLACK
        cell = ws.cell(r, 2, val)
        cell.font = BLUE
        cell.number_format = "mmm d, yyyy" if fmt == "date" else fmt
        ws.cell(r, 3, note).font = NOTE
        N[key] = f"Inputs!$B${r}"
        r += 1
    # month multipliers
    r += 1
    ws.cell(r, 1, "Sales seasonality by month (1.00 = May)").font = BOLD
    r += 1
    mm_first = r
    for m in range(1, 13):
        ws.cell(r, 1, date(2026, m, 1).strftime("%B")).font = BLACK
        c = ws.cell(r, 2, MONTH_MULT[m])
        c.font, c.number_format = BLUE, "0.00"
        r += 1
    MM = f"Inputs!$B${mm_first}:$B${mm_first + 11}"
    N["mm"] = MM
    hist_row = int(N["hist"].split("$")[-1])
    ws.cell(hist_row, 2, f"=AVERAGE(B{mm_first + 6}:B{mm_first + 9})").font = BLACK
    oct_row = int(N["oct_ads"].split("$")[-1])
    ws.cell(oct_row, 2, f"={N['base']}*{N['ad_rate']}*B{mm_first + 9}/{N['hist']}").font = BLACK
    for k in ("hist", "oct_ads"):
        ws.cell(int(N[k].split("$")[-1]), 2).font = BLACK
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60
    for key in ("opening", "buffer"):
        ws[N[key].replace("Inputs!", "").replace("$", "")].fill = YELLOW

    # ------------------------------------------------------------------ Calendar
    cal = wb.create_sheet("Calendar")
    heads = ["Date", "Week", "Month key", "Sales intensity", "Prior month key", "Share of week", "Shopify gross",
             "Card spend", "Card autopay due", "Fixed costs due", "Interest due", "Sales tax due"]
    for i, h in enumerate(heads, 1):
        c = cal.cell(1, i, h)
        c.font, c.fill = BOLD, HEAD
    R = f"$2:${CAL_END_ROW}"
    rng = lambda L: f"Calendar!${L}$2:${L}${CAL_END_ROW}"
    for k in range(CAL_DAYS):
        rr = 2 + k
        d = CAL_START + timedelta(days=k)
        cal.cell(rr, 1, d).number_format = "ddd mmm d, yyyy"
        cal.cell(rr, 1).font = BLUE
        f = {
            2: f"=INT((A{rr}-{N['week0']})/7)",
            3: f"=YEAR(A{rr})*100+MONTH(A{rr})",
            4: (f"=IF(AND(A{rr}>={N['bfcm_start']},A{rr}<={N['bfcm_end']}),INDEX({MM},11)*30*{N['bfcm_share']}/7,"
                f"IF(MONTH(A{rr})=11,INDEX({MM},11)*30*(1-{N['bfcm_share']})/23,INDEX({MM},MONTH(A{rr}))))"),
            5: f"=IF(MONTH(A{rr})=1,(YEAR(A{rr})-1)*100+12,YEAR(A{rr})*100+MONTH(A{rr})-1)",
            6: f"=D{rr}/SUMIFS($D$2:$D${CAL_END_ROW},$B$2:$B${CAL_END_ROW},B{rr})",
            7: f"=INDEX(Forecast!$B$8:$Z$8,1,{pos(f'B{rr}')})*F{rr}",
            8: (f"=IF(B{rr}>=0,(INDEX(Forecast!$B$8:$Z$8,1,{pos(f'B{rr}')})*{N['ad_rate']}+{N['card_other']})/7,"
                f"IF(C{rr}=202610,({N['oct_ads']}+{N['card_other']})/7,0))"),
            9: f"=IF(DAY(A{rr})=22,SUMIFS($H$2:$H${CAL_END_ROW},$C$2:$C${CAL_END_ROW},E{rr}),0)",
            10: f"=IF(DAY(A{rr})=1,{N['fixed']},0)",
            11: f"=IF(DAY(A{rr})=1,{N['loc_drawn']}*{N['apr']}/12,0)",
            12: f"=IF(DAY(A{rr})=20,{N['tax']}*SUMIFS($G$2:$G${CAL_END_ROW},$C$2:$C${CAL_END_ROW},E{rr}),0)",
        }
        for ci, fx in f.items():
            c = cal.cell(rr, ci, fx)
            c.font = BLACK
            c.number_format = USD if ci >= 7 else ("0.000" if ci in (4, 6) else "0")
    cal.freeze_panes = "B2"
    for i, w in enumerate([18, 7, 10, 13, 13, 12, 14, 12, 15, 14, 12, 13], 1):
        cal.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------ Bills
    bl = wb.create_sheet("Bills")
    bh = ["Bill", "Vendor", "Cash line", "Issue week", "Due week", "Pay week", "Priced on month", "Priced on sales",
          "Rate", "Amount", "Early-pay discount", "Take discount (1/0)", "Cash paid", "Discount deadline week",
          "Criticality"]
    for i, h in enumerate(bh, 1):
        c = bl.cell(1, i, h)
        c.font, c.fill = BOLD, HEAD
        c.alignment = Alignment(wrap_text=True)
    week_of_day = lambda y, m, d: f"INT((DATE({y},{m},{d})-{N['week0']})/7)"
    bill_rows = []
    for po in cfg["purchase_orders"]:
        bill_rows.append(dict(id=po["id"], vendor=po["vendor"], line="inventory", issue=po["due_week"] - 2,
                              due=po["due_week"], amount=po["amount"], disc=0, crit="critical"))
    for fr in cfg["freight_invoices"]:
        bill_rows.append(dict(id=fr["id"], vendor=cfg["forwarder"], line="freight_duty", issue=fr["issue_week"],
                              due=fr["issue_week"] + 2, amount=fr["amount"], disc=0, crit="high"))
    for (y, m) in [(2026, 11), (2026, 12), (2027, 1)]:
        py, pm = (y, m - 1) if m > 1 else (y - 1, 12)
        bill_rows.append(dict(id=f"TPL-{y}{m:02d}", vendor=cfg["threepl"], line="fulfillment_3pl",
                              issue=f"={week_of_day(y, m, 10)}", due="issue+4", basis=py * 100 + pm,
                              rate=N["_tpl"] if False else cfg["tpl_pct"], disc=0.02, crit="high"))
    for (y, m) in [(2026, 11), (2026, 12), (2027, 1)]:
        py, pm = (y, m - 1) if m > 1 else (y - 1, 12)
        bill_rows.append(dict(id=f"PKG-{y}{m:02d}", vendor=cfg["packaging_vendor"], line="other_ap",
                              issue=f"={week_of_day(y, m, 5)}", due="issue+4", basis=py * 100 + pm,
                              rate=cfg["packaging_pct"], disc=0.01, crit="normal"))
    bill_rows.append(dict(id="SAAS-ANNUAL", vendor=cfg["saas_vendor"], line="other_ap", issue=5, due=9,
                          amount=cfg["saas_annual"], disc=0, crit="low"))
    for i, b in enumerate(bill_rows):
        rr = 2 + i
        bl.cell(rr, 1, b["id"]).font = BLACK
        bl.cell(rr, 2, b["vendor"]).font = BLACK
        bl.cell(rr, 3, b["line"]).font = BLACK
        bl.cell(rr, 4, b["issue"]).font = BLACK if isinstance(b["issue"], str) else BLUE
        if b["due"] == "issue+4":
            bl.cell(rr, 5, f"=D{rr}+4").font = BLACK
        else:
            bl.cell(rr, 5, b["due"]).font = BLUE
        c = bl.cell(rr, 6, f"=E{rr}")
        c.font, c.fill = BLACK, YELLOW
        c.comment = Comment("Pay week. Defaults to the due week; type a week number to test a pay plan.", "AR")
        if "basis" in b:
            bl.cell(rr, 7, b["basis"]).font = BLUE
            bl.cell(rr, 8, f"=SUMIFS({rng('G')},{rng('C')},G{rr})").font = GREEN
            c = bl.cell(rr, 9, b["rate"])
            c.font, c.number_format = BLUE, PCT
            bl.cell(rr, 10, f"=H{rr}*I{rr}").font = BLACK
        else:
            bl.cell(rr, 10, b["amount"]).font = BLUE
        c = bl.cell(rr, 11, b["disc"])
        c.font, c.number_format = BLUE, PCT
        bl.cell(rr, 12, 0).font = BLUE
        bl.cell(rr, 13, f"=J{rr}*(1-K{rr}*L{rr})").font = BLACK
        bl.cell(rr, 14, f'=IF(K{rr}>0,D{rr}+1,"-")').font = BLACK
        bl.cell(rr, 15, b["crit"]).font = BLACK
        for ci in (8, 10, 13):
            bl.cell(rr, ci).number_format = USD
    B_END = 1 + len(bill_rows)
    bl.cell(B_END + 2, 1, "Set column L to 1 and column F to the discount deadline to take an early-pay "
                          "discount; the Forecast tab updates.").font = NOTE
    for i, w in enumerate([14, 22, 15, 10, 10, 10, 12, 14, 8, 12, 11, 11, 12, 12, 11], 1):
        bl.column_dimensions[get_column_letter(i)].width = w
    bills_sum = lambda line, w: (f"-SUMIFS(Bills!$M$2:$M${B_END},Bills!$C$2:$C${B_END},\"{line}\","
                                 f"Bills!$F$2:$F${B_END},{w})")

    # ------------------------------------------------------------------ Forecast
    fs = wb.create_sheet("Forecast", 1)
    fs["A1"] = "Fernwood Goods: 13-week direct cash forecast (synthetic)"
    fs["A1"].font = TITLE
    fs["A2"] = ("Grey columns are history weeks used by lags and terms. Cash rows start in week 0. Positive = cash "
                "in.")
    fs["A2"].font = NOTE
    weeks = list(range(FIRST_W, LAST_W + 1))
    fwd = [w for w in weeks if w >= 0]
    labels = {4: "Week", 5: "Week starting", 6: "Sales multiplier vs average history week", 8: "Shopify gross sales",
              9: "Refund rate", 10: "Refunds", 11: "Processing fees", 12: "Chargebacks", 13: "Shopify net sales",
              14: "Shopify payouts", 16: "Amazon sales", 17: "Amazon settlements", 19: "Faire invoices issued",
              20: "Hudson Mercantile invoices issued", 21: "Wholesale receipts", 23: "Payroll run this week (1/0)",
              24: "Payroll", 25: "Inventory purchase orders", 26: "Freight and duty", 27: "3PL invoices",
              28: "Packaging and software bills", 29: "Card autopay", 30: "Fixed costs", 31: "Sales tax",
              32: "Credit line interest", 33: "Bank fees and other", 35: "Net cash flow", 36: "Opening cash",
              37: "Closing cash", 38: "Cash buffer", 39: "Under buffer by", 41: "Watch: ad spend charged to the card "
                                                                                 "(paid with next month's autopay)"}
    for rr, t in labels.items():
        fs.cell(rr, 1, t).font = BOLD if rr in (14, 17, 21, 35, 37) else BLACK
    pay_weeks = set(cfg["payroll_weeks"])
    for w in weeks:
        L = col(w)
        hist = w < 0
        fs[f"{L}4"] = w
        fs[f"{L}4"].font = BLUE
        fs[f"{L}5"] = f"={N['week0']}+7*{L}4"
        fs[f"{L}5"].number_format = "mmm d"
        fs[f"{L}6"] = f"=SUMIFS({rng('D')},{rng('B')},{L}$4)/7/{N['hist']}"
        fs[f"{L}6"].number_format = "0.000"
        fs[f"{L}8"] = f"={N['base']}*{L}6"
        fs[f"{L}9"] = f"=IF(MONTH({L}5)=1,{N['refund_jan']},{N['refund']})"
        fs[f"{L}9"].number_format = PCT
        fs[f"{L}10"] = f"={L}8*{L}9"
        fs[f"{L}11"] = f"={L}8*{N['fee']}"
        fs[f"{L}12"] = f"={L}8*{N['cb']}"
        fs[f"{L}13"] = f"={L}8-{L}10-{L}11-{L}12"
        fs[f"{L}16"] = f"={N['amz']}*{L}6"
        fs[f"{L}19"] = f"={N['wh']}*{N['faire_share']}*{L}6"
        fs[f"{L}20"] = f"={N['wh']}*{N['hud_share']}*{L}6"
        if not hist:
            P = col(w - 1)
            P2 = col(w - 2)
            fs[f"{L}14"] = f"=(1-{N['lag']})*{L}13+{N['lag']}*{P}13"
            fs[f"{L}17"] = f"=IF(MOD({L}$4,2)=0,({P2}16+{P}16)*(1-{N['amz_fee']}),0)"
            fs[f"{L}21"] = (f"=INDEX($B$19:$Z$19,1,{pos(f'{L}$4-' + N['faire_terms'])})"
                            f"+INDEX($B$20:$Z$20,1,{pos(f'{L}$4-' + N['hud_terms'])})")
            fs[f"{L}23"] = 1 if w in pay_weeks else 0
            fs[f"{L}23"].font = BLUE
            fs[f"{L}24"] = f"=-{L}23*{N['payroll']}"
            fs[f"{L}25"] = "=" + bills_sum("inventory", f"{L}$4")
            fs[f"{L}26"] = "=" + bills_sum("freight_duty", f"{L}$4")
            fs[f"{L}27"] = "=" + bills_sum("fulfillment_3pl", f"{L}$4")
            fs[f"{L}28"] = "=" + bills_sum("other_ap", f"{L}$4")
            fs[f"{L}29"] = f"=-SUMIFS({rng('I')},{rng('B')},{L}$4)"
            fs[f"{L}30"] = f"=-SUMIFS({rng('J')},{rng('B')},{L}$4)"
            fs[f"{L}31"] = f"=-SUMIFS({rng('L')},{rng('B')},{L}$4)"
            fs[f"{L}32"] = f"=-SUMIFS({rng('K')},{rng('B')},{L}$4)"
            fs[f"{L}33"] = f"=-{N['other_opex']}"
            fs[f"{L}35"] = f"={L}14+{L}17+{L}21+SUM({L}24:{L}33)"
            fs[f"{L}36"] = f"={N['opening']}" if w == 0 else f"={P}37"
            fs[f"{L}37"] = f"={L}36+{L}35"
            fs[f"{L}38"] = f"={N['buffer']}"
            fs[f"{L}39"] = f"=MAX(0,{L}38-{L}37)"
            fs[f"{L}41"] = f"=-{L}8*{N['ad_rate']}"
        for rr in range(4, 42):
            c = fs[f"{L}{rr}"]
            if c.value is None:
                continue
            if c.font != BLUE and rr != 4:
                c.font = BLACK
            if rr >= 8 and rr not in (9, 23):
                c.number_format = USD
            if hist:
                c.fill = GREY
        fs[f"{L}37"].font = BOLD
        fs[f"{L}35"].border = THIN
    first, last = col(0), col(LAST_W)
    s = 44
    summary = [("Low point (closing cash)", f"=MIN({first}37:{last}37)", USD),
               ("Low point week", f"=INDEX({first}5:{last}5,1,MATCH(B{s},{first}37:{last}37,0))", "mmm d, yyyy"),
               ("Gap to buffer", f"=MAX(0,{N['buffer']}-B{s})", USD),
               ("Credit draw to hold the buffer", f"=CEILING(B{s + 2},{N['draw_inc']})", USD),
               ("Credit headroom", f"={N['loc_limit']}-{N['loc_drawn']}", USD),
               ("Draw fits the line?", f'=IF(B{s + 3}<=B{s + 4},"yes","no")', "@"),
               ("Ending cash", f"={last}37", USD)]
    fs.cell(s - 1, 1, "Summary").font = BOLD
    for i, (lab, fx, fmt) in enumerate(summary):
        fs.cell(s + i, 1, lab).font = BLACK
        c = fs.cell(s + i, 2, fx)
        c.font, c.number_format = BOLD, fmt
    fs.column_dimensions["A"].width = 46
    for w in weeks:
        fs.column_dimensions[col(w)].width = 11
    fs.freeze_panes = "B6"

    # ------------------------------------------------------------------ Check
    ck = wb.create_sheet("Check")
    ck["A1"] = "Workbook vs Python engine (world/engine.py), same inputs"
    ck["A1"].font = TITLE
    ck["A2"] = "Column C is pasted from a Python run; column B is this workbook's formula result."
    ck["A2"].font = NOTE
    for i, h in enumerate(["Week starting", "Workbook closing cash", "Python closing cash", "Difference",
                           "Difference %"], 1):
        c = ck.cell(4, i, h)
        c.font, c.fill = BOLD, HEAD
    for w in fwd:
        rr = 5 + w
        ck.cell(rr, 1, f"=Forecast!{col(w)}5").number_format = "mmm d"
        ck.cell(rr, 2, f"=Forecast!{col(w)}37").number_format = USD
        c = ck.cell(rr, 3, round(float(fc.cash[w]), 2))
        c.font, c.number_format = BLUE, USD
        ck.cell(rr, 4, f"=B{rr}-C{rr}").number_format = USD
        ck.cell(rr, 5, f"=IF(C{rr}=0,0,D{rr}/C{rr})").number_format = "0.00%"
    end = 5 + LAST_W
    ck.cell(end + 2, 1, "Largest difference %").font = BOLD
    c = ck.cell(end + 2, 2, f"=MAX(MAX(E5:E{end}),-MIN(E5:E{end}))")
    c.font, c.number_format = BOLD, "0.000%"
    ck.cell(end + 3, 1, "Python low point").font = BLACK
    c = ck.cell(end + 3, 2, round(fc.low_point()[1], 2))
    c.font, c.number_format = BLUE, USD
    ck.cell(end + 4, 1, "Workbook low point").font = BLACK
    c = ck.cell(end + 4, 2, "=Forecast!B44")
    c.font, c.number_format = GREEN, USD
    for L, w in zip("ABCDE", (16, 22, 20, 14, 14)):
        ck.column_dimensions[L].width = w
    for wsx in wb.worksheets:
        for row in wsx.iter_rows():
            for c in row:
                if c.font and c.font.name != F:
                    c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic,
                                  color=c.font.color)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


def recalc():
    """Let LibreOffice compute every formula so the file opens with values (formulas are kept)."""
    import shutil
    import subprocess
    import tempfile
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        print("LibreOffice not found: open the workbook in Excel to calculate it")
        return False
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([exe, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(OUT)],
                       check=True, capture_output=True, timeout=180)
        shutil.move(str(Path(tmp) / OUT.name), OUT)
    print("recalculated with LibreOffice")
    return True


if __name__ == "__main__":
    build()
    recalc()
